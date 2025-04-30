import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

# Initialize session state
if "step" not in st.session_state:
    st.session_state.step = 1

# Define poultry process steps
poultry_step_options = {
    "Slaughter": "Bleeding and feather removal",
    "Evisceration": "Internal organ removal",
    "Chilling": "Rapid cooling to prevent pathogen growth",
    "Packaging": "Wrapping and sealing",
    "Labeling": "Label application and handling instructions"
}

# Function to generate Word document

def generate_word_doc(results):
    doc = Document()

    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_run = header_para.add_run("HACCP Hazard Analysis Report")
    header_run.font.name = 'Arial'
    header_run.font.size = Pt(24)
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    steps = {}
    ccp_items = []
    for item in results:
        step = item['Step']
        if step not in steps:
            steps[step] = []
        steps[step].append(item)
        if "CCP" in item["CCP Decision"]:
            ccp_items.append(item)

    for item in ccp_items:
        doc.add_page_break()
        heading = doc.add_heading(f"CCP Detail – {item['Hazard Type']} Hazard at {item['Step']}", level=2)
        heading.runs[0].font.bold = True
        heading.runs[0].font.size = Pt(14)

        ccp_table = doc.add_table(rows=10, cols=2)
        ccp_table.style = 'Table Grid'
        ccp_table.cell(0, 0).merge(ccp_table.cell(0, 1))
        header_cell = ccp_table.cell(0, 0).paragraphs[0].add_run(f"{item['Hazard Type']} Hazard – CCP Documentation")
        header_cell.bold = True
        header_cell.font.size = Pt(12)

        labels = [
            "Step", "Critical Limits", "Monitoring Procedures", "Monitoring Frequency",
            "Monitoring Personnel", "Recordkeeping", "Corrective Actions", "Verification Procedures", "Justification"
        ]
        keys = labels

        for i in range(1, 10):
            ccp_table.cell(i, 0).text = labels[i - 1]
            ccp_table.cell(i, 1).text = item.get(keys[i - 1], "")

    for step, items in steps.items():
        doc.add_page_break()
        heading = doc.add_heading(f"Step: {step}", level=2)
        heading.runs[0].font.bold = True
        heading.runs[0].font.size = Pt(14)

        table = doc.add_table(rows=4, cols=5)
        table.style = 'Table Grid'
        table.cell(0, 0).text = "Hazard Type"
        table.cell(1, 0).text = "Biological"
        table.cell(2, 0).text = "Chemical"
        table.cell(3, 0).text = "Physical"

        headers = ["RLTO?", "Justification", "Control Measures", "Is this Step a CCP"]
        for i, h in enumerate(headers):
            run = table.cell(0, i+1).paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(11)

        for row_idx, hazard in enumerate(["Biological", "Chemical", "Physical"], start=1):
            match = next((i for i in items if i["Hazard Type"] == hazard), None)
            if match:
                table.cell(row_idx, 1).text = "Yes" if match["RLTO"] else "No"
                table.cell(row_idx, 2).text = match.get("Justification", "")
                table.cell(row_idx, 3).text = match.get("CCP Decision", "")
                table.cell(row_idx, 4).text = "Yes" if "CCP" in match["CCP Decision"] else "No"

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Streamlit UI
if st.session_state.step == 1:
    st.title("Step 1: Select Poultry Processing Steps")
    selected_steps = []
    for step, desc in poultry_step_options.items():
        if st.checkbox(f"{step} — {desc}", key=f"{step}_select"):
            selected_steps.append(step)

    if st.button("Next"):
        st.session_state.selected_steps = selected_steps
        st.session_state.step = 2
        st.rerun()

elif st.session_state.step == 2:
    st.title("Step 2: Hazard Analysis")
    st.write(f"Selected steps: {', '.join(st.session_state.selected_steps)}")

    results = []

    def ccp_questions(step, hazard):
        st.markdown("### 🧪 CCP Required Elements (9 CFR § 417.2(c))")
        return {
            "Critical Limits": st.text_input("Critical Limits", key=f"{step}_{hazard}_cl"),
            "Monitoring Procedures": st.text_input("Monitoring Procedures", key=f"{step}_{hazard}_mp"),
            "Monitoring Frequency": st.text_input("Monitoring Frequency", key=f"{step}_{hazard}_mf"),
            "Monitoring Personnel": st.text_input("Monitoring Personnel", key=f"{step}_{hazard}_person"),
            "Recordkeeping": st.text_input("Recordkeeping", key=f"{step}_{hazard}_rk"),
            "Corrective Actions": st.text_input("Corrective Actions", key=f"{step}_{hazard}_ca"),
            "Verification Procedures": st.text_input("Verification Procedures", key=f"{step}_{hazard}_vp")
        }

    for step in st.session_state.selected_steps:
        st.header(f"🔹 Step: {step}")
        for hazard in ["Biological", "Chemical", "Physical"]:
            st.subheader(f"⚠️ {hazard} Hazard")
            rlto = st.radio(f"Is this {hazard.lower()} hazard reasonably likely to occur (RLTO)?", ["Yes", "No"], key=f"{step}_{hazard}_rlto")

            if rlto == "Yes":
                preventive = st.radio("Are preventive measures in place?", ["Yes", "No"], key=f"{step}_{hazard}_preventive")

                if preventive == "No":
                    sop_type = "SOP" if hazard == "Biological" else "SOP/SSOP"
                    controlled = st.radio(f"Is the hazard controlled by a validated {sop_type}?", ["Yes", "No"], key=f"{step}_{hazard}_sop")
                    if controlled == "Yes":
                        justification = st.text_area(
                            "📌 Based on your responses, this hazard is not classified as a CCP. Please explain why:",
                            key=f"{step}_{hazard}_sop_justification"
                        )
                        results.append({
                            "Step": step,
                            "Hazard Type": hazard,
                            "RLTO": True,
                            "CCP Decision": f"Not a CCP – Controlled by validated {sop_type}",
                            "Justification": justification
                        })
                        continue

                contam = st.radio("Could contamination occur at this step?", ["Yes", "No"], key=f"{step}_{hazard}_contam")
                if contam == "No":
                    justification = st.text_area("📌 Based on your responses, this hazard is not classified as a CCP. Please explain:", key=f"{step}_{hazard}_notccp_justification")
                    results.append({
                        "Step": step,
                        "Hazard Type": hazard,
                        "RLTO": True,
                        "CCP Decision": "Not a CCP – Contamination not likely",
                        "Justification": justification
                    })
                    continue

                eliminates = st.radio("Does this step eliminate or reduce the hazard?", ["Yes", "No"], key=f"{step}_{hazard}_eliminates")
                if eliminates == "Yes":
                    sop_type = "SOP" if hazard == "Biological" else "SOP/SSOP"
                    controlled = st.radio(f"Is this control validated by a {sop_type}?", ["Yes", "No"], key=f"{step}_{hazard}_eliminates_sop")
                    if controlled == "Yes":
                        justification = st.text_area("📌 Based on your responses, this hazard is not classified as a CCP. Please explain:", key=f"{step}_{hazard}_eliminated_sop_justification")
                        results.append({
                            "Step": step,
                            "Hazard Type": hazard,
                            "RLTO": True,
                            "CCP Decision": f"Not a CCP – Eliminated and validated by {sop_type}",
                            "Justification": justification
                        })
                        continue

                future = st.radio("Is there a later step that controls this hazard?", ["Yes", "No"], key=f"{step}_{hazard}_future")
                if future == "Yes":
                    justification = st.text_area("📌 Based on your responses, this hazard is not classified as a CCP. Please explain:", key=f"{step}_{hazard}_future_justification")
                    results.append({
                        "Step": step,
                        "Hazard Type": hazard,
                        "RLTO": True,
                        "CCP Decision": "Not a CCP – Controlled later",
                        "Justification": justification
                    })
                    continue

                justification = st.text_area("📌 This hazard is a CCP because no validated control or later step exists. Please explain:", key=f"{step}_{hazard}_ccp_final_justification")
                ccp_data = ccp_questions(step, hazard)
                results.append({
                    "Step": step,
                    "Hazard Type": hazard,
                    "RLTO": True,
                    "CCP Decision": "CCP – No validated control or future step",
                    "Justification": justification,
                    **ccp_data
                })

            else:
                justification = st.text_area("📌 Based on your responses, this hazard is not classified as a CCP. Please explain:", key=f"{step}_{hazard}_notrlto_justification")
                results.append({
                    "Step": step,
                    "Hazard Type": hazard,
                    "RLTO": False,
                    "CCP Decision": "Not a CCP – Hazard not RLTO",
                    "Justification": justification
                })

    if st.button("📄 Generate HACCP Summary"):
        word_file = generate_word_doc(results)
        st.download_button(
            label="⬇️ Download HACCP Summary (.docx)",
            data=word_file,
            file_name="haccp_summary.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
